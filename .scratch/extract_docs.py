"""Extract both DOCX files to readable text."""
from docx import Document
from pathlib import Path

def extract(path):
    doc = Document(path)
    out = []
    for p in doc.paragraphs:
        t = p.text.strip()
        if not t:
            continue
        s = p.style.name if p.style else ''
        if s.startswith('Heading'):
            try:
                n = int(s.replace('Heading ', ''))
                out.append(f"\n{'#' * n} {t}\n")
                continue
            except Exception:
                pass
        out.append(t)
    for table in doc.tables:
        out.append("\n--- TABLE ---")
        for row in table.rows:
            out.append(" | ".join(c.text.strip() for c in row.cells))
    return "\n".join(out)

base = Path('/Users/fivefriday/.hermes/profiles/heidi/cache/documents')
stick_path = base / 'doc_71e95d8287f5_STICK_Brand_Bible_1.docx'
workbook_path = base / 'doc_eb46302b0ffa_Campaign_OS_Full_Brand_Bible_Workbook_2.docx'

print(f'=== STICK Brand Bible ({stick_path.stat().st_size:,} bytes) ===\n')
print(extract(stick_path))

print('\n\n\n=== Campaign OS Full Brand Bible Workbook 2 ===\n')
print(extract(workbook_path))
